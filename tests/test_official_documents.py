from datetime import date
import hashlib
from io import BytesIO

import pytest
from docx import Document
from pypdf import PdfReader

from services.official_documents import fill_land_agreement, fill_stock_control
from services.official_documents import (
    AUTHORSHIP_DECLARATION,
    LAND_AGREEMENT_TEMPLATE,
    STOCK_CONTROL_TEMPLATE,
)


def _data():
    return {
        "landowner": "Anna Beispiel, 9000 St. Gallen",
        "beekeeper": "Joel Beispiel",
        "beekeeper_number": "CH-123",
        "beekeeper_street": "Feldstrasse 1",
        "beekeeper_city": "9000 St. Gallen",
        "phone": "+41 79 000 00 00",
        "email": "joel@example.ch",
        "section": "St. Gallen",
        "parcel": "Parzelle 123, Appenzell",
        "area_m2": 120,
        "installation": "Other",
        "other_installation": "Vier Magazinbeuten",
        "site_plan_attached": True,
        "start_date": date(2026, 4, 15),
        "fixed_term": True,
        "end_date": date(2026, 9, 30),
        "notice_period": "6 months",
        "notice_timing": "At month end",
        "other_notice_rule": "Nach der Honigernte",
        "additional_duty": "Zugangsweg freihalten",
        "compensation": "CHF 100",
        "compensation_period": "Year",
        "veterinary_office": "Amt für Verbraucherschutz und Veterinärwesen",
        "bee_inspector": "Max Muster",
        "apiary_number": "SG-456",
        "site_street": "Parzelle 123",
        "site_city": "9050 Appenzell",
        "coordinates": "47.3310/9.4090",
        "move_date": date(2026, 4, 15),
        "origin_apiary_number": "SG-111",
        "colonies": 8,
        "year": 2026,
    }


def test_land_agreement_keeps_form_and_fills_all_guided_choices():
    reader = PdfReader(BytesIO(fill_land_agreement(_data())))
    fields = reader.get_fields()

    assert len(reader.pages) == 2
    assert len(fields) == 30
    assert fields["Vereinbarung zwischen"]["/V"].startswith("Anna Beispiel")
    assert fields["Imkerin  Imker"]["/V"] == "Joel Beispiel, Feldstrasse 1, 9000 St. Gallen"
    assert fields["Check Box1"]["/V"] == "/Ja"
    assert fields["Ja empfohlen"]["/V"] == "/On"
    assert fields["befristet bis"]["/V"] == "/On"
    assert fields["6 Monate"]["/V"] == "/On"
    assert fields["auf ein Monatsende"]["/V"] == "/On"
    assert fields["Check Box2"]["/V"] == "/Ja"
    assert fields["Jahr"]["/V"] == "/On"
    # The two signature widgets intentionally remain empty.
    assert fields["undefined_5"].get("/V") is None
    assert fields["undefined_6"].get("/V") is None


def test_stock_control_keeps_official_tables_and_prefills_first_move():
    document = Document(BytesIO(fill_stock_control(_data())))

    assert len(document.tables) == 3
    assert document.paragraphs[0].text.endswith("2026")
    assert document.tables[0].cell(4, 1).text == "Joel Beispiel"
    assert document.tables[0].cell(6, 5).text == "47.3310/9.4090"
    movement = document.tables[1].rows[2].cells
    assert movement[0].text == "15.4.26"
    assert movement[1].text == "SG-111"
    assert movement[3].text == "Verstellen"
    assert movement[4].text == "8"


def test_generators_never_modify_the_bundled_official_templates():
    pdf_before = LAND_AGREEMENT_TEMPLATE.read_bytes()
    docx_before = STOCK_CONTROL_TEMPLATE.read_bytes()

    fill_land_agreement(_data())
    fill_stock_control(_data())

    assert LAND_AGREEMENT_TEMPLATE.read_bytes() == pdf_before
    assert STOCK_CONTROL_TEMPLATE.read_bytes() == docx_before


@pytest.mark.parametrize(
    ("choice", "field", "value"),
    [
        ("Bee house", "ein Bienenhaus gestellt Fahrnisbaute Pläne oder Fotos liegen bei", "/On"),
        ("Hives without a bee house", "kein Bienenhaus gestellt Pläne oder Fotos von Bienenbeute liegen bei", "/On"),
        ("Other", "Check Box1", "/Ja"),
    ],
)
def test_every_installation_option_maps_to_the_original_pdf(choice, field, value):
    data = _data()
    data["installation"] = choice
    fields = PdfReader(BytesIO(fill_land_agreement(data))).get_fields()
    assert fields[field]["/V"] == value


@pytest.mark.parametrize(("attached", "field"), [(True, "Ja empfohlen"), (False, "Nein")])
def test_both_site_plan_options_map_to_the_original_pdf(attached, field):
    data = _data()
    data["site_plan_attached"] = attached
    fields = PdfReader(BytesIO(fill_land_agreement(data))).get_fields()
    assert fields[field]["/V"] == "/On"


@pytest.mark.parametrize(
    ("fixed", "field"), [(True, "befristet bis"), (False, "unbefristet")]
)
def test_both_duration_options_map_to_the_original_pdf(fixed, field):
    data = _data()
    data["fixed_term"] = fixed
    fields = PdfReader(BytesIO(fill_land_agreement(data))).get_fields()
    assert fields[field]["/V"] == "/On"


@pytest.mark.parametrize(
    ("choice", "field"),
    [
        ("3 months", "3 Monate"),
        ("6 months", "6 Monate"),
        ("9 months", "9 Monate"),
        ("12 months", "12 Monate"),
    ],
)
def test_every_notice_period_maps_to_the_original_pdf(choice, field):
    data = _data()
    data["notice_period"] = choice
    fields = PdfReader(BytesIO(fill_land_agreement(data))).get_fields()
    assert fields[field]["/V"] == "/On"


@pytest.mark.parametrize(
    ("choice", "field"),
    [
        ("Any time", "beliebig"),
        ("At month end", "auf ein Monatsende"),
        ("Only on a specified date", "nur per"),
    ],
)
def test_every_notice_timing_maps_to_the_original_pdf(choice, field):
    data = _data()
    data["notice_timing"] = choice
    fields = PdfReader(BytesIO(fill_land_agreement(data))).get_fields()
    assert fields[field]["/V"] == "/On"


@pytest.mark.parametrize(("choice", "field"), [("Year", "Jahr"), ("Month", "Monat")])
def test_both_compensation_periods_map_to_the_original_pdf(choice, field):
    data = _data()
    data["compensation_period"] = choice
    fields = PdfReader(BytesIO(fill_land_agreement(data))).get_fields()
    assert fields[field]["/V"] == "/On"


def test_not_selected_contract_options_leave_every_corresponding_box_open():
    data = _data()
    data.update(
        {
            "notice_period": "Not selected",
            "notice_timing": "Not selected",
            "other_notice_rule": "",
            "compensation_period": "Not selected",
        }
    )
    fields = PdfReader(BytesIO(fill_land_agreement(data))).get_fields()
    untouched = (
        "3 Monate",
        "6 Monate",
        "9 Monate",
        "12 Monate",
        "beliebig",
        "auf ein Monatsende",
        "nur per",
        "Check Box2",
        "Jahr",
        "Monat",
    )
    assert all(fields[name].get("/V") is None for name in untouched)


def test_filled_pdf_keeps_logical_values_and_widget_appearances():
    reader = PdfReader(BytesIO(fill_land_agreement(_data())))
    for annotation in reader.pages[0].get("/Annots") or []:
        widget = annotation.get_object()
        if widget.get("/Subtype") != "/Widget" or widget.get("/V") is None:
            continue
        assert widget.get("/AP", {}).get("/N")


def test_hsg_declaration_is_the_unchanged_supplied_official_pdf():
    assert AUTHORSHIP_DECLARATION.exists()
    assert len(PdfReader(AUTHORSHIP_DECLARATION).pages) == 1
    assert hashlib.sha256(AUTHORSHIP_DECLARATION.read_bytes()).hexdigest() == (
        "a257cb09939dd0dbf52625bb2559c137dca740e01b839d130a3869c9213195ff"
    )
